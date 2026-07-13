"""
Generates '288-IFRS18-Demand-Charter.pptx', reproducing the company's
"Demand Charter" slide template layout and filling it in for demand
288 - IFRS 18.

Usage:
    pip install python-pptx
    python3 generate_charter_pptx.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- colors -----------------------------------------------------------
DARK_BLUE = RGBColor(0x1F, 0x3B, 0x73)
HEADER_BLUE = RGBColor(0x8E, 0xA9, 0xDB)
LIGHT_GRAY = RGBColor(0xCF, 0xD3, 0xDA)
ORANGE = RGBColor(0xF4, 0xB1, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    shape.line.width = Pt(0.75)


def add_box(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_text(shape, text, size=11, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font_name="Calibri"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name


def add_bullets(shape, items, size=11, color=BLACK, font_name="Calibri"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"\u2022 {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = font_name
        p.space_after = Pt(4)


prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

# ---- Title --------------------------------------------------------------
title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(10), Inches(0.6))
set_text(title_box, "Demand Charter - 288 - IFRS 18", size=28, bold=True, color=DARK_BLUE)

# ---- Header info row (Titel / Demand Manager / E2E-Core Proc / Security) --
header_y = Inches(0.85)
header_h = Inches(0.4)
labels = [
    ("Titel", 1.1, HEADER_BLUE, WHITE),
    ("288 - IFRS 18", 2.6, LIGHT_GRAY, BLACK),
    ("Demand\nManager", 1.3, HEADER_BLUE, WHITE),
    ("Enik\u0151 Baga", 2.6, LIGHT_GRAY, BLACK),
    ("E2E /\nCore Proc.", 1.3, HEADER_BLUE, WHITE),
    ("Finance \u2013 Accounting &\nFinancial Reporting (R2R)", 2.7, LIGHT_GRAY, BLACK),
    ("Security\nClassif.", 1.1, HEADER_BLUE, WHITE),
    ("Internal", 1.4, LIGHT_GRAY, BLACK),
]
x = Inches(0.3)
for text, width_in, fill, fontcolor in labels:
    w = Inches(width_in)
    box = add_box(slide, x, header_y, w, header_h, fill_color=fill)
    bold = fill == HEADER_BLUE
    size = 10 if fill == HEADER_BLUE else 10
    set_text(box, text, size=size, bold=bold, color=fontcolor,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x += w

# ---- Two main sections: Current Situation / Business Objectives ---------
section_y = Inches(1.35)
section_header_h = Inches(0.32)
section_body_h = Inches(2.3)
col_w = Inches(6.35)
gap = Inches(0.1)

left1 = Inches(0.3)
left2 = left1 + col_w + gap

h1 = add_box(slide, left1, section_y, col_w, section_header_h, fill_color=HEADER_BLUE)
set_text(h1, "Current Situation and Pain Points", size=12, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

h2 = add_box(slide, left2, section_y, col_w, section_header_h, fill_color=HEADER_BLUE)
set_text(h2, "Business Objectives and High-Level Demand Description", size=12, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

body1 = add_box(slide, left1, section_y + section_header_h, col_w, section_body_h,
                 fill_color=LIGHT_GRAY)
add_bullets(body1, [
    "IAS 1 today does not mandate P&L categories/subtotals \u2192 inconsistent "
    "presentation across entities and limited peer comparability.",
    "No formal governance for management-defined performance measures (MPMs) "
    "used in investor communications \u2192 audit/compliance risk.",
    "Chart of accounts and consolidation/reporting system not structured to "
    "auto-classify P&L lines into operating / investing / financing.",
    "Alternative performance measures produced via manual spreadsheets with "
    "limited audit trail.",
    "Inconsistent aggregation/disaggregation of line items across business units.",
    "No dedicated budget/resourcing yet assigned for IFRS 18 readiness.",
], size=10.5)

body2 = add_box(slide, left2, section_y + section_header_h, col_w, section_body_h,
                 fill_color=LIGHT_GRAY)
add_bullets(body2, [
    "Achieve full IFRS 18 compliance before mandatory effective date "
    "(periods beginning \u2265 1 Jan 2027; comparatives from 1 Jan 2026).",
    "Gap assessment of chart of accounts, consolidation system and "
    "disclosure templates vs. the 3 new categories & 2 new subtotals.",
    "Define & govern Management-Defined Performance Measures (MPMs), "
    "incl. required reconciliation disclosures.",
    "Reconfigure P&L structure, CoA mapping and consolidation tool for "
    "native category/subtotal reporting.",
    "Update statutory templates, disclosure checklists and cash-flow "
    "reconciliation logic; train Finance/Controlling/IR.",
    "Run a parallel/dry-run cycle and align with external auditors on "
    "transition approach.",
], size=10.5)

# ---- Benefits & Stakeholders & Allocations banner ------------------------
banner_y = section_y + section_header_h + section_body_h + Inches(0.08)
banner_h = Inches(0.3)
banner = add_box(slide, left1, banner_y, col_w * 2 + gap, banner_h, fill_color=HEADER_BLUE)
set_text(banner, "Benefits & Stakeholders & Allocations", size=12, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---- 4 columns: Qualitative / Quantitative / Stakeholders / Allocations --
col4_y = banner_y + banner_h
col4_header_h = Inches(0.3)
col4_body_h = Inches(2.05)
col4_w = Inches(3.1625)

col4_titles = ["Qualitative Benefits", "Quantitative Benefits", "Stakeholders", "Allocations Demand"]
col4_bullets = [
    [
        "Full compliance with mandatory standard; avoids audit qualification.",
        "Improved comparability/transparency for investors via standardized subtotals.",
        "Stronger governance/audit trail over non-GAAP measures.",
        "Harmonized P&L supports statutory & management reporting.",
        "Reduced manual effort/error risk at close.",
    ],
    [
        "Reduced manual closing/adjustment effort per cycle once automated "
        "(FTE-days TBD after design phase).",
        "Avoidance of one-off remediation costs (advisory, late changes, "
        "audit findings) from a rushed transition.",
        "Precise figures to be confirmed after gap assessment.",
    ],
    [
        "Group Finance / Corporate Accounting (owner)",
        "Group Controlling / FP&A",
        "Investor Relations",
        "Consolidation & Reporting Systems / IT",
        "Internal Audit & External Auditors",
        "Business Unit Controllers",
    ],
    None,  # handled separately below
]

x = left1
for i in range(4):
    header = add_box(slide, x, col4_y, col4_w, col4_header_h, fill_color=ORANGE)
    set_text(header, col4_titles[i], size=11, bold=True, color=BLACK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    body = add_box(slide, x, col4_y + col4_header_h, col4_w, col4_body_h, fill_color=LIGHT_GRAY)
    if i < 3:
        add_bullets(body, col4_bullets[i], size=9.5)
    else:
        set_text(
            body,
            "Demand effort CIT (PD):\n~40 PD (illustrative; IT/consolidation "
            "system config, CoA mapping, reports)\n\n"
            "Demand effort BG (PD):\n~25 PD (illustrative; gap assessment, "
            "MPM governance, policy, training)",
            size=9.5,
        )
    x += col4_w + Emu(0)

# ---- Footer ---------------------------------------------------------------
footer = slide.shapes.add_textbox(Inches(9.8), Inches(7.1), Inches(3.3), Inches(0.3))
set_text(footer, "INNOVATING TOGETHER", size=11, bold=True, color=DARK_BLUE,
         align=PP_ALIGN.RIGHT)

out_path = "288-IFRS18-Demand-Charter.pptx"
prs.save(out_path)
print(f"Saved {out_path}")
