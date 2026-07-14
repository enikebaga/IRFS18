"""
Generates '288-IFRS18-Demand-Charter.docx' — a Word / knowledge-base version
of the filled-in Demand Charter for demand 288 - IFRS 18.

Usage:
    pip install python-docx
    python3 generate_charter_docx.py

Run this from within demand-charter/; the output (.docx) belongs
alongside the .md/.pptx versions of the same charter in this folder.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK_BLUE = RGBColor(0x1F, 0x3B, 0x73)
HEADER_BLUE = RGBColor(0x8E, 0xA9, 0xDB)
LIGHT_GRAY = RGBColor(0xCF, 0xD3, 0xDA)
ORANGE = RGBColor(0xF4, 0xB1, 0x83)


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            if align is not None:
                p.alignment = align
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def add_bullets(cell, items, size=9.5):
    cell.text = ""
    for i, item in enumerate(items):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = None
        run = p.add_run(f"\u2022 {item}")
        run.font.size = Pt(size)


doc = Document()

# base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

section = doc.sections[0]
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

title = doc.add_heading(level=0)
run = title.add_run("Demand Charter - 288 - IFRS 18")
run.font.color.rgb = DARK_BLUE
run.font.size = Pt(24)

doc.add_paragraph(
    "Knowledge-base copy of the filled-in Demand Charter for demand 288 "
    "(IFRS 18 Implementation), written in plain, easy-to-understand "
    "language. See also the .md and .pptx versions in this folder; all "
    "three contain the same content."
).italic = True

# ---- Header info table ----------------------------------------------------
header_table = doc.add_table(rows=2, cols=4)
header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
header_labels = ["Titel", "Demand Manager", "E2E / Core Proc.", "Security Classif."]
header_values = [
    "288 - IFRS 18",
    "Enik\u0151 Baga",
    "Finance \u2013 Accounting & Financial Reporting",
    "Internal",
]
for i, (label, value) in enumerate(zip(header_labels, header_values)):
    hcell = header_table.rows[0].cells[i]
    set_cell_text(hcell, label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(hcell, "8EA9DB")
    vcell = header_table.rows[1].cells[i]
    set_cell_text(vcell, value, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(vcell, "CFD3DA")

doc.add_paragraph()

# ---- Current Situation / Business Objectives -------------------------------
section_table = doc.add_table(rows=2, cols=2)
section_table.alignment = WD_TABLE_ALIGNMENT.CENTER

h1 = section_table.rows[0].cells[0]
set_cell_text(h1, "What's the problem today?\n(Current Situation and Pain Points)", bold=True,
              color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
shade_cell(h1, "8EA9DB")

h2 = section_table.rows[0].cells[1]
set_cell_text(h2, "What are we trying to achieve?\n(Business Objectives and High-Level Demand Description)", bold=True,
              color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
shade_cell(h2, "8EA9DB")

pain_points = [
    "No fixed structure today, so different teams present profit & loss "
    "differently \u2192 hard to compare across the company or with peers.",
    "No clear, approved process for our \u201cadjusted\u201d performance numbers "
    "shown to investors \u2192 risk once we must explain how they're calculated.",
    "Our systems don't automatically sort P&L items into the 3 new required "
    "groups: Operating, Investing, Financing.",
    "Adjusted performance numbers are built by hand in spreadsheets \u2014 "
    "slow, error-prone, and hard for auditors to check.",
    "Business units group similar costs differently; the new rule needs "
    "much more consistency here.",
    "No dedicated time or budget yet assigned to prepare for this change.",
    "We haven't yet talked to our external auditors about how we'll make "
    "this change or restate last year's numbers.",
]
b1 = section_table.rows[1].cells[0]
add_bullets(b1, pain_points)
shade_cell(b1, "CFD3DA")

objectives = [
    "Be fully ready for IFRS 18 before it's mandatory (years starting "
    "on/after 1 Jan 2027; last year's figures restated from 1 Jan 2026).",
    "Check what needs to change in our chart of accounts, reporting system "
    "and templates for the 3 new groups & 2 new required subtotals.",
    "Set clear, approved rules for our \u201cadjusted\u201d performance numbers, "
    "incl. showing how they connect back to the official numbers.",
    "Update our P&L structure, account mapping and reporting system so the "
    "new groups/subtotals are produced automatically.",
    "Update official templates, checklists and cash-flow statement; train "
    "Finance/Controlling/Investor Relations.",
    "Run a test/dry-run reporting cycle and work with our external auditors "
    "on how we make the change.",
]
b2 = section_table.rows[1].cells[1]
add_bullets(b2, objectives)
shade_cell(b2, "CFD3DA")

doc.add_paragraph()

# ---- Benefits & Stakeholders & Allocations banner --------------------------
# A 1x1 shaded table gives a solid colored banner, matching the section
# headers above it.
banner_tbl = doc.add_table(rows=1, cols=1)
bcell = banner_tbl.rows[0].cells[0]
set_cell_text(bcell, "Benefits & Stakeholders & Allocations", bold=True,
              color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
shade_cell(bcell, "8EA9DB")

# ---- 4 columns table -------------------------------------------------------
four_table = doc.add_table(rows=2, cols=4)
four_table.alignment = WD_TABLE_ALIGNMENT.CENTER

titles = ["Qualitative Benefits", "Quantitative Benefits", "Stakeholders", "Allocations Demand"]
for i, t in enumerate(titles):
    c = four_table.rows[0].cells[i]
    set_cell_text(c, t, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c, "F4B183")

qualitative = [
    "We follow the new mandatory rule, avoiding audit and regulatory risk.",
    "Our financial statements become easier to compare, for our teams and investors.",
    "Clearer, better-documented rules for our adjusted performance numbers.",
    "Simpler, more consistent P&L for both official and internal reporting.",
    "Less manual work and fewer mistakes at closing time.",
]
quantitative = [
    "Less manual work each closing cycle once the system sorts things "
    "automatically (exact savings TBD after detailed review).",
    "We avoid extra costs (advisors, rushed fixes, audit findings) from "
    "leaving this until the last minute.",
    "Exact figures to be confirmed after the initial review.",
]
stakeholders = [
    "Group Finance / Corporate Accounting (owner)",
    "Group Controlling / FP&A",
    "Investor Relations",
    "IT / Reporting Systems team",
    "Internal Audit & External Auditors",
    "Local Finance teams (Business Units)",
    "Legal / Company Secretary (if applicable)",
]

c0 = four_table.rows[1].cells[0]
add_bullets(c0, qualitative)
shade_cell(c0, "CFD3DA")

c1 = four_table.rows[1].cells[1]
add_bullets(c1, quantitative)
shade_cell(c1, "CFD3DA")

c2 = four_table.rows[1].cells[2]
add_bullets(c2, stakeholders)
shade_cell(c2, "CFD3DA")

c3 = four_table.rows[1].cells[3]
set_cell_text(
    c3,
    "Demand effort CIT (PD):\n~40 person-days (rough estimate; updating "
    "IT systems, accounts, reports)\n\n"
    "Demand effort BG (PD):\n~25 person-days (rough estimate; reviewing "
    "changes, setting rules, training)",
    size=9.5,
)
shade_cell(c3, "CFD3DA")

doc.add_paragraph()
note = doc.add_paragraph()
note_run = note.add_run(
    "Note: This is a first draft prepared to unblock kickoff. The effort "
    "estimates under \u201cAllocations Demand\u201d and the quantitative benefits are "
    "placeholders and should be reviewed and confirmed by the Demand Manager "
    "and the relevant IT/Finance leads before formal approval."
)
note_run.italic = True
note_run.font.size = Pt(9)

footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer_p.add_run("INNOVATING TOGETHER")
footer_run.bold = True
footer_run.font.color.rgb = DARK_BLUE

out_path = "288-IFRS18-Demand-Charter.docx"
doc.save(out_path)
print(f"Saved {out_path}")
